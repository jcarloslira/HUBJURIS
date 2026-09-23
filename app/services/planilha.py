"""Planilhas Excel sob demanda — o que o escritório manda para a administradora.

Relatório em PDF é para o síndico ler. Planilha é para o escritório trabalhar:
conferir unidade a unidade com a lista da administradora, somar, filtrar, mandar
para o contador. Até aqui só existia o menu "Exportar", com as colunas que nós
escolhemos; o que o Dr. Wilker pediu foi poder dizer em português o que quer na
planilha e recebê-la pronta.

Duas decisões que separam isto de um CSV renomeado:

- **as somas são fórmulas** (``=SUM(D2:D31)``), não números congelados: mexeu numa
  linha, o total acompanha, que é como uma planilha tem de se comportar;
- **cada coluna tem tipo** — dinheiro sai como moeda (soma no Excel, alinha à
  direita), data como data, percentual como percentual. Texto "R$ 1.234,56" numa
  célula é um número perdido.
"""

import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from typing import Any

from openpyxl import Workbook
from openpyxl.drawing.image import Image
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.services.documentos import Timbre
from app.services.pdf_modelo import cor_primaria

# Mesma paleta do PDF: a planilha é do mesmo escritório.
_PETROLEO = "2C5254"
_DOURADO = "A79E6E"
_LINHA_PAR = "F4F6F5"
_CINZA = "DFE3E1"
_CINZA_TEXTO = "5E6B6D"
# Altura do logo na faixa, em pixels: grande o bastante para se reconhecer a
# marca, pequeno o bastante para não empurrar a tabela para fora da primeira tela.
_LOGO_ALTURA = 44


def _cor_da_marca(cor: str) -> str:
    """A cor do escritório em ARGB do Excel, com o mesmo critério do PDF.

    ``cor_primaria`` devolve o petróleo quando a cor cadastrada é clara demais
    para o texto branco do cabeçalho: sem isso, um escritório que cadastrou o
    dourado do logo ganharia uma tabela ilegível.
    """
    return cor_primaria(cor).lstrip("#").upper()

FORMATOS = {
    "texto": "@",
    "numero": "#,##0",
    "dinheiro": 'R$ #,##0.00',
    "data": "DD/MM/YYYY",
    "percentual": "0.0%",
}
TIPOS_SOMAVEIS = {"numero", "dinheiro"}

_LARGURA_MINIMA = 10
_LARGURA_MAXIMA = 62
_MAX_ABAS = 8
_MAX_LINHAS = 5_000
_MAX_COLUNAS = 30


class PlanilhaError(Exception):
    """Pedido de planilha que não dá para montar."""


@dataclass
class Coluna:
    """Uma coluna: o título que o usuário lê e o tipo que o Excel entende."""

    titulo: str
    tipo: str = "texto"

    def __post_init__(self) -> None:
        self.titulo = " ".join(str(self.titulo or "").split()) or "Coluna"
        self.tipo = self.tipo if self.tipo in FORMATOS else "texto"


@dataclass
class Aba:
    """Uma aba da planilha."""

    nome: str
    colunas: list[Coluna]
    linhas: list[list[Any]] = field(default_factory=list)
    somar: list[str] = field(default_factory=list)
    nota: str = ""

    def __post_init__(self) -> None:
        # O Excel recusa : \ / ? * [ ] no nome da aba e corta em 31 caracteres.
        limpo = re.sub(r"[:\\/?*\[\]]", " ", str(self.nome or "Dados"))
        self.nome = " ".join(limpo.split())[:31] or "Dados"


def _numero(texto: str) -> float | None:
    """Lê o número escrito em português OU em inglês, sem perder os centavos.

    O modelo escreve das duas formas na mesma planilha: "R$ 26.651,45" e
    "26651.45". Tratar o ponto sempre como milhar multiplicava o valor por cem, e
    numa planilha que vai para a administradora o erro passa por plausível. A
    regra é a do separador que aparece por ÚLTIMO: ele é o decimal.
    """
    limpo = re.sub(r"[^\d,.\-]", "", texto)
    if not limpo:
        return None
    ponto, virgula = limpo.rfind("."), limpo.rfind(",")
    if ponto >= 0 and virgula >= 0:
        decimal, milhar = (".", ",") if ponto > virgula else (",", ".")
    elif virgula >= 0:
        decimal, milhar = ",", "."  # português: a vírgula decide
    elif ponto >= 0:
        casas = len(limpo.split(".")[-1])
        # "1.234.567" e "R$ 1.500" são milhar; "26651.45" e "0.5" são decimais.
        milhar_provavel = limpo.count(".") > 1 or casas == 3
        decimal, milhar = ("", ".") if milhar_provavel else (".", ",")
    else:
        decimal, milhar = "", ""
    if milhar:
        limpo = limpo.replace(milhar, "")
    if decimal and decimal != ".":
        limpo = limpo.replace(decimal, ".")
    try:
        return float(limpo)
    except ValueError:
        return None


def _valor(bruto: Any, tipo: str) -> Any:
    """Converte o que o modelo escreveu para o que o Excel guarda de verdade."""
    if bruto is None or bruto == "":
        return None
    if tipo in ("numero", "dinheiro", "percentual"):
        if isinstance(bruto, (int, float)):
            return float(bruto) / 100 if tipo == "percentual" and abs(bruto) > 1 else float(bruto)
        texto = str(bruto).strip()
        numero = _numero(texto)
        if numero is None:
            return str(bruto)
        if texto.startswith("(") and texto.endswith(")"):  # contábil: (1.234) é negativo
            numero = -abs(numero)
        return numero / 100 if tipo == "percentual" and abs(numero) > 1 else numero
    if tipo == "data":
        if isinstance(bruto, (date, datetime)):
            return bruto
        texto = str(bruto).strip()[:10]
        for formato in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
            try:
                return datetime.strptime(texto, formato).date()
            except ValueError:
                continue
        return str(bruto)
    return str(bruto)


def _marca(planilha: Worksheet, aba: Aba, timbre: Timbre | None, cor: str) -> int:
    """Escreve a faixa de identidade no topo. Devolve a linha do cabeçalho.

    A planilha sai do escritório para a administradora e para o contador: chegar
    com o nome e o logo de quem mandou é o mínimo. Sem timbre, começa direto na
    tabela, sem faixa vazia.
    """
    nome = (timbre.nome if timbre else "").strip()
    logo = timbre.logo if timbre else None
    if not nome and not logo:
        return 1

    linha = 1
    if logo is not None:
        try:
            imagem = Image(io.BytesIO(logo))
            escala = min(_LOGO_ALTURA / (imagem.height or 1), 1)
            imagem.height = int((imagem.height or 1) * escala)
            imagem.width = int((imagem.width or 1) * escala)
            imagem.anchor = "A1"
            planilha.add_image(imagem)
            planilha.row_dimensions[1].height = _LOGO_ALTURA * 0.78
            linha = 2
        except Exception:  # noqa: BLE001 - logo quebrado não impede a planilha
            linha = 1

    ultima_coluna = get_column_letter(len(aba.colunas))
    if nome:
        planilha.merge_cells(f"A{linha}:{ultima_coluna}{linha}")
        celula = planilha.cell(row=linha, column=1, value=nome)
        celula.font = Font(bold=True, size=13, color=cor)
        celula.alignment = Alignment(vertical="center")
        planilha.row_dimensions[linha].height = 22
        linha += 1

    subtitulo = " · ".join(p for p in (aba.nome, aba.nota) if p)
    if subtitulo:
        planilha.merge_cells(f"A{linha}:{ultima_coluna}{linha}")
        celula = planilha.cell(row=linha, column=1, value=subtitulo)
        celula.font = Font(size=9, color=_CINZA_TEXTO)
        planilha.row_dimensions[linha].height = 15
        linha += 1

    # Filete dourado separando a marca da tabela, como na capa do PDF.
    for coluna in range(1, len(aba.colunas) + 1):
        planilha.cell(row=linha, column=coluna).border = Border(
            bottom=Side(style="medium", color=_DOURADO)
        )
    planilha.row_dimensions[linha].height = 6
    return linha + 1


def _escrever_aba(planilha: Worksheet, aba: Aba, timbre: Timbre | None = None) -> None:
    cor = _cor_da_marca(timbre.cor if timbre else "")
    topo = _marca(planilha, aba, timbre, cor)
    cabecalho = Font(bold=True, color="FFFFFF", size=10)
    fundo = PatternFill("solid", fgColor=cor)
    zebra = PatternFill("solid", fgColor=_LINHA_PAR)
    borda = Border(bottom=Side(style="thin", color=_CINZA))

    for coluna, definicao in enumerate(aba.colunas, start=1):
        celula = planilha.cell(row=topo, column=coluna, value=definicao.titulo)
        celula.font = cabecalho
        celula.fill = fundo
        celula.alignment = Alignment(vertical="center", wrap_text=True)
    planilha.row_dimensions[topo].height = 26

    for indice, linha in enumerate(aba.linhas[:_MAX_LINHAS]):
        numero = topo + 1 + indice
        for coluna, definicao in enumerate(aba.colunas, start=1):
            bruto = linha[coluna - 1] if coluna - 1 < len(linha) else None
            celula = planilha.cell(row=numero, column=coluna, value=_valor(bruto, definicao.tipo))
            celula.number_format = FORMATOS[definicao.tipo]
            celula.border = borda
            celula.alignment = Alignment(
                vertical="top",
                horizontal="right" if definicao.tipo in TIPOS_SOMAVEIS else "left",
                wrap_text=definicao.tipo == "texto",
            )
            if indice % 2:
                celula.fill = zebra

    ultima = topo + len(aba.linhas[:_MAX_LINHAS])
    _totalizar(planilha, aba, ultima, topo, cor)
    _dimensionar(planilha, aba)
    planilha.freeze_panes = f"A{topo + 1}"
    if ultima > topo:
        planilha.auto_filter.ref = f"A{topo}:{get_column_letter(len(aba.colunas))}{ultima}"


def _totalizar(planilha: Worksheet, aba: Aba, ultima: int, topo: int, cor: str) -> None:
    """Linha de total com fórmula viva — mexeu na tabela, o total acompanha."""
    somaveis = {t.lower() for t in aba.somar}
    alvos = [
        i
        for i, c in enumerate(aba.colunas, start=1)
        if c.tipo in TIPOS_SOMAVEIS and (not somaveis or c.titulo.lower() in somaveis)
    ]
    if not alvos or ultima <= topo:
        return
    total = ultima + 1
    primeira = topo + 1
    rotulo = planilha.cell(row=total, column=1, value="TOTAL")
    rotulo.font = Font(bold=True, color=cor)
    for coluna in alvos:
        letra = get_column_letter(coluna)
        celula = planilha.cell(
            row=total, column=coluna, value=f"=SUM({letra}{primeira}:{letra}{ultima})"
        )
        celula.number_format = FORMATOS[aba.colunas[coluna - 1].tipo]
        celula.font = Font(bold=True, color=cor)
        celula.alignment = Alignment(horizontal="right")
    for coluna in range(1, len(aba.colunas) + 1):
        planilha.cell(row=total, column=coluna).border = Border(
            top=Side(style="medium", color=_DOURADO)
        )


def _dimensionar(planilha: Worksheet, aba: Aba) -> None:
    """Largura pelo conteúdo: coluna estreita esconde o dado, larga demais rola a tela."""
    for coluna, definicao in enumerate(aba.colunas, start=1):
        tamanhos = [len(definicao.titulo)]
        for linha in aba.linhas[:200]:
            if coluna - 1 < len(linha):
                tamanhos.append(len(str(linha[coluna - 1] or "")))
        largura = max(tamanhos) + 3
        planilha.column_dimensions[get_column_letter(coluna)].width = max(
            _LARGURA_MINIMA, min(_LARGURA_MAXIMA, largura)
        )


def montar_planilha(
    abas: list[Aba], *, autor: str = "", timbre: Timbre | None = None
) -> bytes:
    """Monta o .xlsx. Levanta PlanilhaError quando não há o que escrever."""
    reais = [a for a in abas if a.colunas][:_MAX_ABAS]
    if not reais:
        raise PlanilhaError("Nenhuma aba com colunas para montar.")
    livro = Workbook()
    livro.remove(livro.active)
    usados: set[str] = set()
    for aba in reais:
        nome = aba.nome
        sufixo = 2
        while nome.lower() in usados:  # duas abas com o mesmo nome quebram o arquivo
            nome = f"{aba.nome[:28]} {sufixo}"
            sufixo += 1
        usados.add(nome.lower())
        _escrever_aba(livro.create_sheet(title=nome), aba, timbre)
    if autor:
        livro.properties.creator = autor
    buffer = io.BytesIO()
    livro.save(buffer)
    return buffer.getvalue()


def abas_do_modelo(dados: Any) -> list[Aba]:
    """Traduz o JSON da ferramenta em abas, descartando em silêncio o que vier torto."""
    bruto = dados.get("abas") if isinstance(dados, dict) else dados
    if isinstance(bruto, dict):
        bruto = [bruto]
    if not isinstance(bruto, list):
        return []
    abas: list[Aba] = []
    for item in bruto[:_MAX_ABAS]:
        if not isinstance(item, dict):
            continue
        colunas = [
            Coluna(c.get("titulo", ""), str(c.get("tipo", "texto")).lower())
            if isinstance(c, dict)
            else Coluna(str(c))
            for c in (item.get("colunas") or [])[:_MAX_COLUNAS]
        ]
        if not colunas:
            continue
        linhas = [
            list(linha)[:_MAX_COLUNAS]
            for linha in (item.get("linhas") or [])
            if isinstance(linha, (list, tuple))
        ]
        somar = [str(s) for s in (item.get("somar") or []) if str(s).strip()]
        abas.append(
            Aba(
                nome=str(item.get("nome") or "Dados"),
                colunas=colunas,
                linhas=linhas,
                somar=somar,
                nota=str(item.get("nota") or ""),
            )
        )
    return abas
