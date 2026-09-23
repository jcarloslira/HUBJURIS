"""Planilha Excel sob demanda.

O que se testa aqui é o que separa uma planilha de um CSV renomeado: valor em
reais tem de chegar como NÚMERO (senão não soma no Excel), data como data, e o
total tem de ser fórmula viva.
"""

import io

import pytest
from openpyxl import load_workbook

from app.services.planilha import (
    Aba,
    Coluna,
    PlanilhaError,
    abas_do_modelo,
    montar_planilha,
)

PEDIDO = {
    "abas": [
        {
            "nome": "Unidades em cobrança",
            "colunas": [
                {"titulo": "Unidade", "tipo": "texto"},
                {"titulo": "Distribuição", "tipo": "data"},
                {"titulo": "Valor da causa", "tipo": "dinheiro"},
            ],
            "linhas": [
                ["204-B", "12/03/2025", "R$ 15.462,33"],
                ["1403-A", "2025-06-01", 28450.7],
                ["802", "05/01/2026", "R$ 1.234.567,89"],
            ],
            "somar": ["Valor da causa"],
        }
    ]
}


def _abrir(dados: bytes):
    return load_workbook(io.BytesIO(dados))


def test_valor_em_reais_chega_como_numero_e_soma() -> None:
    """Texto "R$ 15.462,33" numa célula é um número perdido: não soma, não ordena."""
    planilha = _abrir(montar_planilha(abas_do_modelo(PEDIDO))).active

    assert planilha.cell(row=2, column=3).value == pytest.approx(15462.33)
    assert planilha.cell(row=4, column=3).value == pytest.approx(1234567.89)


def test_data_chega_como_data_nos_dois_formatos_que_o_modelo_escreve() -> None:
    planilha = _abrir(montar_planilha(abas_do_modelo(PEDIDO))).active

    assert planilha.cell(row=2, column=2).value.strftime("%d/%m/%Y") == "12/03/2025"
    assert planilha.cell(row=3, column=2).value.strftime("%d/%m/%Y") == "01/06/2025"


def test_total_e_formula_viva_nao_numero_congelado() -> None:
    """Corrigiu uma linha na mão? O total tem de acompanhar, como em qualquer planilha."""
    planilha = _abrir(montar_planilha(abas_do_modelo(PEDIDO))).active

    assert planilha.cell(row=5, column=3).value == "=SUM(C2:C4)"
    assert planilha.cell(row=5, column=1).value == "TOTAL"


def test_planilha_ja_vem_utilizavel_com_filtro_e_cabecalho_fixo() -> None:
    planilha = _abrir(montar_planilha(abas_do_modelo(PEDIDO))).active

    assert planilha.freeze_panes == "A2"
    assert planilha.auto_filter.ref == "A1:C4"


def test_coluna_de_texto_nao_ganha_linha_de_total() -> None:
    dados = montar_planilha(
        [Aba("Lista", [Coluna("Assunto"), Coluna("Responsável")], [["Cobrança", "Ana"]])]
    )
    planilha = _abrir(dados).active

    assert planilha.cell(row=3, column=1).value is None


def test_somar_escolhe_a_coluna_certa_quando_ha_duas_numericas() -> None:
    dados = montar_planilha(
        [
            Aba(
                "Carteira",
                [Coluna("Unidade"), Coluna("Ações", "numero"), Coluna("Valor", "dinheiro")],
                [["204-B", 2, 1000.0], ["802", 1, 500.0]],
                somar=["Valor"],
            )
        ]
    )
    planilha = _abrir(dados).active

    assert planilha.cell(row=4, column=2).value is None  # "Ações" ficou de fora
    assert planilha.cell(row=4, column=3).value == "=SUM(C2:C3)"


def test_valor_que_nao_e_numero_nao_quebra_a_planilha() -> None:
    """O modelo escreve "a apurar" numa célula de valor: a planilha sai assim mesmo."""
    dados = montar_planilha(
        [Aba("Carteira", [Coluna("Valor", "dinheiro")], [["a apurar"], [1500.0]])]
    )
    planilha = _abrir(dados).active

    assert planilha.cell(row=2, column=1).value == "a apurar"
    assert planilha.cell(row=3, column=1).value == pytest.approx(1500.0)


def test_nome_de_aba_com_caractere_proibido_e_saneado() -> None:
    """':' e '/' no nome da aba fazem o Excel recusar o arquivo inteiro."""
    aba = Aba("Processos: 2025/2026", [Coluna("Unidade")], [["204-B"]])

    assert ":" not in aba.nome and "/" not in aba.nome


def test_duas_abas_com_o_mesmo_nome_nao_quebram_o_arquivo() -> None:
    dados = montar_planilha(
        [
            Aba("Resumo", [Coluna("A")], [["x"]]),
            Aba("Resumo", [Coluna("B")], [["y"]]),
        ]
    )

    assert len(_abrir(dados).sheetnames) == 2


def test_linha_mais_curta_que_o_cabecalho_nao_derruba_a_geracao() -> None:
    dados = montar_planilha(
        [Aba("Lista", [Coluna("Unidade"), Coluna("Valor", "dinheiro")], [["204-B"]])]
    )
    planilha = _abrir(dados).active

    assert planilha.cell(row=2, column=2).value is None


def test_pedido_sem_coluna_vira_erro_de_dominio() -> None:
    with pytest.raises(PlanilhaError):
        montar_planilha([])


def test_json_torto_do_modelo_e_descartado_em_silencio() -> None:
    abas = abas_do_modelo(
        {
            "abas": [
                {"nome": "Boa", "colunas": [{"titulo": "X"}], "linhas": [["1"], "lixo"]},
                {"nome": "Sem colunas", "linhas": [["1"]]},
                "isto não é aba",
            ]
        }
    )

    assert [a.nome for a in abas] == ["Boa"]
    assert abas[0].linhas == [["1"]]
